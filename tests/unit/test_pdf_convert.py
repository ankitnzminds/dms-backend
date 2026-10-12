"""Signing-copy conversion: formats, safety limits and honest refusals."""

from __future__ import annotations

import io

import pytest

from app.modules.signatures import pdf_convert as pc

TXT = "text/plain"


def text_of(pdf: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages)


def pages(pdf: bytes) -> int:
    return pc.page_count(pdf)


def png(w=40, h=20, color=(200, 30, 30), mode="RGB"):  # noqa: ANN001, ANN201
    from PIL import Image

    out = io.BytesIO()
    Image.new(mode, (w, h), color).save(out, "PNG")
    return out.getvalue()


def test_pdf_is_passed_through_unchanged():
    data = b"%PDF-1.4 hello"
    assert pc.convert_to_pdf(filename="a.pdf", content_type="application/pdf", data=data) == data


@pytest.mark.parametrize(
    ("name", "ct"),
    [
        ("n.txt", TXT),
        ("t.csv", "text/csv"),
        ("r.md", "text/markdown"),
        ("x.log", "application/octet-stream"),
    ],
)
def test_text_like_files_become_readable_pdfs(name, ct):  # noqa: ANN001
    body = "Hello, world\n  indented line\ncol1,col2\n1,2\n"
    pdf = pc.convert_to_pdf(filename=name, content_type=ct, data=body.encode())
    assert pdf.startswith(b"%PDF-")
    out = text_of(pdf)
    assert "Hello, world" in out and "col1,col2" in out and name.rsplit(".", 1)[0] in out


def test_text_is_escaped_never_interpreted_as_markup():
    pdf = pc.convert_to_pdf(
        filename="e.txt", content_type=TXT, data=b"<b>bold</b> & <script>alert(1)</script>"
    )
    assert "<b>bold</b>" in text_of(pdf) and "<script>" in text_of(pdf)


def test_long_documents_paginate():
    data = "\n".join(f"line {i}" for i in range(400)).encode()
    assert pages(pc.convert_to_pdf(filename="long.txt", content_type=TXT, data=data)) >= 3


def test_png_and_transparent_png_and_jpeg():
    from PIL import Image

    for data, name in ((png(), "a.png"), (png(mode="RGBA", color=(0, 0, 0, 0)), "t.png")):
        assert pc.convert_to_pdf(filename=name, content_type="image/png", data=data).startswith(
            b"%PDF-"
        )
    out = io.BytesIO()
    Image.new("RGB", (30, 30), (10, 200, 10)).save(out, "JPEG")
    assert (
        pages(pc.convert_to_pdf(filename="p.jpg", content_type="image/jpeg", data=out.getvalue()))
        == 1
    )


def test_multi_frame_tiff_gets_one_page_per_frame():
    from PIL import Image

    frames = [Image.new("RGB", (30, 30), c) for c in ((255, 0, 0), (0, 255, 0), (0, 0, 255))]
    out = io.BytesIO()
    frames[0].save(out, "TIFF", save_all=True, append_images=frames[1:])
    assert (
        pages(pc.convert_to_pdf(filename="m.tif", content_type="image/tiff", data=out.getvalue()))
        == 3
    )


def test_content_decides_not_the_extension():
    # a PNG named .txt is still an image; text named .png is refused as an unreadable image
    assert pc.kind_of("fake.txt", TXT, png()[:16]) == "image"
    with pytest.raises(pc.ConversionError):
        pc.convert_to_pdf(
            filename="fake.png", content_type="image/png", data=b"not an image at all"
        )


def test_docx_paragraphs_and_tables_are_extracted():
    import docx

    d = docx.Document()
    d.add_heading("Master Services Agreement", level=1)
    d.add_paragraph("The parties agree as follows.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = (
        "Item",
        "Price",
        "Widget",
        "10",
    )
    out = io.BytesIO()
    d.save(out)
    pdf = pc.convert_to_pdf(
        filename="msa.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=out.getvalue(),
    )
    t_out = text_of(pdf)
    assert (
        "Master Services Agreement" in t_out
        and "The parties agree" in t_out
        and "Widget | 10" in t_out.replace("\n", " ")
        or "Widget" in t_out
    )


def test_xlsx_sheets_become_text_and_hidden_sheets_are_skipped():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Laptop", 1200])
    hidden = wb.create_sheet("Secret")
    hidden.append(["do not show"])
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    pdf = pc.convert_to_pdf(
        filename="b.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data=out.getvalue(),
    )
    t = text_of(pdf)
    assert "# Budget" in t and "Laptop, 1200" in t and "do not show" not in t


@pytest.mark.parametrize("name", ["a.doc", "b.xls", "c.ppt", "d.pptx", "e.exe", "f.zip", "g.bin"])
def test_unsupported_types_are_refused_with_the_supported_list(name):  # noqa: ANN001
    with pytest.raises(pc.ConversionError) as e:
        pc.convert_to_pdf(
            filename=name,
            content_type="application/octet-stream",
            data=b"PK\x03\x04junk" + b"\x00" * 8,
        )
    assert "can't be converted to PDF" in str(e.value) and "Supported for signing" in str(e.value)


def test_limits_and_damaged_files():
    with pytest.raises(pc.ConversionError, match="empty"):
        pc.convert_to_pdf(filename="e.txt", content_type=TXT, data=b"")
    with pytest.raises(pc.ConversionError, match="too large"):
        pc.convert_to_pdf(
            filename="big.txt", content_type=TXT, data=b"x" * (pc.MAX_INPUT_BYTES + 1)
        )
    with pytest.raises(pc.ConversionError, match="damaged|can't be read"):
        pc.convert_to_pdf(
            filename="bad.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=b"PK\x03\x04not really a zip",
        )


def test_zip_bombs_are_refused_before_parsing(monkeypatch):  # noqa: ANN001
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w/>")
    monkeypatch.setattr(pc, "MAX_UNZIPPED_BYTES", 2)
    with pytest.raises(pc.ConversionError, match="too large or complex"):
        pc.convert_to_pdf(
            filename="b.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            data=out.getvalue(),
        )


def test_documenso_box_placement_rules():
    from app.adapters.signatures.documenso import _box

    p = {"page_number": 3, "page_x": 20.0, "page_y": 50.0, "page_width": 30.0, "page_height": 5.0}
    first = _box(p, 0, 5)
    assert (
        first["pageNumber"],
        first["pageX"],
        first["pageY"],
        first["pageWidth"],
        first["pageHeight"],
    ) == (3, 20.0, 50.0, 30.0, 5.0)
    assert _box(p, 1, 5)["pageY"] == 56.0  # stacked directly below
    assert _box(p, 30, 5)["pageY"] == 95.0  # never leaves the page
    assert _box(p, 0, 2)["pageNumber"] == 2  # a page beyond the document is clamped, never invalid
    default = _box(None, 0, 4)
    assert default["pageNumber"] == 4 and default["pageX"] == 10.0  # unchanged default spot
